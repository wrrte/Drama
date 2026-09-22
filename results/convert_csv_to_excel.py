import argparse
import ast
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# (Random, Human, DramaXS 논문) 고정 참조 점수.
# Random/Human은 STORM/results/convert_csv_to_excel.py와 같은 정밀도의 기준값,
# DramaXS는 사용자가 첨부한 논문 표의 DramaXS 열입니다.
REFERENCE_SCORES = {
    "Alien": (227.8, 7127.7, 820),
    "Amidar": (5.8, 1719.5, 131),
    "Assault": (222.4, 742.0, 539),
    "Asterix": (210.0, 8503.3, 1632),
    "BankHeist": (14.2, 753.1, 137),
    "BattleZone": (2360.0, 37187.5, 10860),
    "Boxing": (0.1, 12.1, 78),
    "Breakout": (1.7, 30.5, 7),
    "ChopperCommand": (811.0, 7387.8, 1642),
    "CrazyClimber": (10780.5, 35829.4, 83931),
    "DemonAttack": (152.1, 1971.0, 201),
    "Freeway": (0.0, 29.6, 15),
    "Frostbite": (65.2, 4334.7, 785),
    "Gopher": (257.6, 2412.5, 2757),
    "Hero": (1027.0, 30826.4, 7946),
    "Jamesbond": (29.0, 302.8, 372),
    "Kangaroo": (52.0, 3035.0, 1384),
    "Krull": (1598.0, 2665.5, 9693),
    "KungFuMaster": (258.5, 22736.3, 23920),
    "MsPacman": (307.3, 6951.6, 2270),
    "Pong": (-20.7, 14.6, 15),
    "PrivateEye": (24.9, 69571.3, 90),
    "Qbert": (163.9, 13455.0, 796),
    "RoadRunner": (11.5, 7845.0, 14020),
    "Seaquest": (68.4, 42054.7, 497),
    "UpNDown": (533.4, 11693.2, 7387),
}

PAPER_SCORE_COLUMN = "DramaXS (논문 점수)"
PAIRED_MEAN_COLUMN = "Mean (공통 시드)"
SCORE_DELTA_COLUMN = "Δ Score (행별 비교)"
HNS_DELTA_COLUMN = "Δ HNS (행별 비교)"


def load_excluded_seeds():
    """update_tex.py를 실행하지 않고 현재 EXCLUDED_SEEDS 설정을 읽습니다."""
    source_path = Path(__file__).resolve().with_name("update_tex.py")
    tree = ast.parse(source_path.read_text(encoding="utf-8"), filename=str(source_path))
    for node in tree.body:
        if isinstance(node, ast.Assign):
            targets = node.targets
        elif isinstance(node, ast.AnnAssign):
            targets = [node.target]
        else:
            continue
        if any(isinstance(target, ast.Name) and target.id == "EXCLUDED_SEEDS" for target in targets):
            return ast.literal_eval(node.value)
    raise ValueError(f"EXCLUDED_SEEDS 설정을 찾을 수 없습니다: {source_path}")


def mark_excluded_seeds(worksheet, results, seed_columns, excluded_seeds):
    """점수를 보존하면서 제외된 시드의 X/O 셀을 회색과 취소선으로 표시합니다."""
    excluded_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
    for row_index, game in enumerate(results["Game"], start=2):
        for seed in seed_columns:
            if int(seed) not in excluded_seeds.get(game, set()):
                continue
            cell = worksheet.cell(row=row_index, column=results.columns.get_loc(seed) + 1)
            font = copy(cell.font)
            font.strike = True
            if "RUNNING" not in str(cell.value):
                cell.fill = excluded_fill
                font.color = "808080"
            cell.font = font
            cell.comment = Comment(
                f"EXCLUDED_SEEDS: {game}, seed {seed}\n"
                "Drama/results/update_tex.py의 EXCLUDED_SEEDS에 지정되어 "
                "공통 시드 평균, Δ Score, Δ HNS 및 LaTeX 결과 집계에서 제외되는 시드입니다.\n"
                "취소선은 제외된 시드, 노란색 배경은 실행 중인 run을 뜻합니다.",
                "Drama",
            )


def add_paired_comparisons(results, seed_columns, excluded_seeds):
    """제외 목록을 반영한 X/O 공통 시드 평균과 논문/target 16 비교를 추가합니다."""
    results = results.copy()
    results[" "] = ""
    for column in (PAPER_SCORE_COLUMN, PAIRED_MEAN_COLUMN, SCORE_DELTA_COLUMN, HNS_DELTA_COLUMN):
        results[column] = np.nan

    for game, group in results.groupby("Game"):
        baseline_rows = group.index[group["Retrieval"].eq("X")]
        target_rows = group.index[group["Retrieval"].eq("O")]
        if baseline_rows.empty:
            continue
        baseline_index = baseline_rows[0]
        reference = REFERENCE_SCORES.get(game)
        if reference is not None:
            results.at[baseline_index, PAPER_SCORE_COLUMN] = reference[2]
        if target_rows.empty:
            continue
        target_index = target_rows[0]
        game_seed_columns = [
            seed for seed in seed_columns
            if int(seed) not in excluded_seeds.get(game, set())
        ]
        baseline = pd.to_numeric(results.loc[baseline_index, game_seed_columns], errors="coerce")
        target = pd.to_numeric(results.loc[target_index, game_seed_columns], errors="coerce")
        valid = baseline.notna() & target.notna()
        if not valid.any():
            continue

        baseline_mean = baseline[valid].mean()
        target_mean = target[valid].mean()
        target_delta = target_mean - baseline_mean
        results.at[baseline_index, PAIRED_MEAN_COLUMN] = baseline_mean
        results.at[target_index, PAIRED_MEAN_COLUMN] = target_mean
        # STORM과 동일하게 target 행의 Δ Score는 절댓값, Δ HNS는 부호를 유지합니다.
        results.at[target_index, SCORE_DELTA_COLUMN] = abs(target_delta)
        if reference is not None:
            random_score, human_score, paper_score = reference
            paper_delta = baseline_mean - paper_score
            results.at[baseline_index, SCORE_DELTA_COLUMN] = paper_delta
            if human_score != random_score:
                results.at[baseline_index, HNS_DELTA_COLUMN] = paper_delta / (human_score - random_score)
                results.at[target_index, HNS_DELTA_COLUMN] = target_delta / (human_score - random_score)
    return results


def format_score(value):
    if pd.isna(value) or str(value).strip() in {"", "N/A", "nan"}:
        return ""
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return str(value)


def main():
    parser = argparse.ArgumentParser(description="Convert Drama WandB CSV to an O/X comparison workbook.")
    parser.add_argument("--input", default="drama_wandb_runs.csv")
    parser.add_argument("--output", default="drama_results.xlsx")
    args = parser.parse_args()
    excluded_seeds = load_excluded_seeds()

    runs = pd.read_csv(args.input)
    if runs.empty:
        raise ValueError("The input CSV contains no O/X runs.")

    runs["Created At"] = pd.to_datetime(runs["Created At"], utc=True, errors="coerce").dt.tz_localize(None)
    runs["State"] = runs["State"].astype(str).str.strip().str.lower()
    runs["Eval Return Numeric"] = pd.to_numeric(runs["Eval Return"], errors="coerce")
    runs = runs[runs["Retrieval"].isin(["O", "X", "BOTH"])].copy()
    if runs.empty:
        raise ValueError("The input CSV contains no O/X runs.")

    running_both = runs[
        (runs["Retrieval"] == "BOTH")
        & (runs["State"].astype(str).str.lower() == "running")
    ].copy()
    if not running_both.empty:
        running_both = pd.concat(
            [running_both.assign(Retrieval="X"), running_both.assign(Retrieval="O")],
            ignore_index=True,
        )
    runs = runs[runs["Retrieval"].isin(["O", "X"])].copy()
    runs = pd.concat([runs, running_both], ignore_index=True)

    # O행과 오른쪽 비교는 target 16 실험을 사용합니다. X는 target 설정과 무관합니다.
    target_16 = pd.to_numeric(runs["Retrieval Target"], errors="coerce").eq(16)
    runs = runs[runs["Retrieval"].eq("X") | target_16].copy()
    if runs.empty:
        raise ValueError("The input CSV contains no retrieval-off or target-16 runs.")

    # Keep the newest run when the same game, seed, and mode was logged more than once.
    runs = runs.sort_values("Created At").drop_duplicates(
        ["Game", "Seed", "Retrieval"], keep="last"
    )
    seed_columns = sorted(runs["Seed"].dropna().unique())
    results = runs.pivot_table(
        index=["Game", "Retrieval"],
        columns="Seed",
        values="Eval Return Numeric",
        aggfunc="first",
    ).reindex(columns=seed_columns).reindex(pd.MultiIndex.from_product(
        [sorted(runs["Game"].unique()), ["X", "O"]],
        names=["Game", "Retrieval"],
    ))
    results.columns.name = None
    results = results.reset_index()
    results = results[["Game", "Retrieval"] + sorted(seed_columns)]

    running = runs[runs["State"].astype(str).str.lower() == "running"].pivot_table(
        index=["Game", "Retrieval"],
        columns="Seed",
        values="State",
        aggfunc="first",
    ).reindex(index=results.set_index(["Game", "Retrieval"]).index, columns=seed_columns)
    running.columns.name = None
    running = running.reset_index(drop=True)

    detail = runs.drop(columns=["Eval Return Numeric"])
    detail["Eval Return"] = detail["Eval Return"].map(format_score)
    detail["Eval Normalized Return"] = detail["Eval Normalized Return"].map(format_score)
    for column in results.columns[2:]:
        results[column] = results[column].map(format_score)
        # 중간 평가 점수가 있어도 RUNNING 셀은 공통 시드 평균에서 제외합니다.
        results.loc[running[column].eq("running"), column] = "RUNNING"
    results = add_paired_comparisons(results, seed_columns, excluded_seeds)

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        results.to_excel(writer, sheet_name="Results", index=False)
        detail.to_excel(writer, sheet_name="Runs", index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.font = Font(name="Calibri", size=12, bold=cell.row == 1)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_cells in worksheet.columns:
                width = max(len(str(cell.value or "")) for cell in column_cells)
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 12), 36)

        running_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        results_sheet = writer.book["Results"]
        results_sheet.column_dimensions[get_column_letter(len(seed_columns) + 3)].width = 3
        summary_formats = {
            PAPER_SCORE_COLUMN: (30, "0"),
            PAIRED_MEAN_COLUMN: (30, "0.00"),
            SCORE_DELTA_COLUMN: (46, "+0.00;-0.00;0.00"),
            HNS_DELTA_COLUMN: (32, "+0.0000;-0.0000;0.0000"),
        }
        for name, (width, number_format) in summary_formats.items():
            column = results.columns.get_loc(name) + 1
            results_sheet.column_dimensions[get_column_letter(column)].width = width
            for row in results_sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                row[0].number_format = (
                    "0.00" if name == SCORE_DELTA_COLUMN
                    and results_sheet.cell(row=row[0].row, column=2).value == "O"
                    else number_format
                )
        comments = {
            PAPER_SCORE_COLUMN: "첨부된 논문 표의 DramaXS 게임별 점수입니다. X행에 표시합니다.",
            PAIRED_MEAN_COLUMN: (
                "X: Retrieval 미사용, O: Retrieval 적용 (target: 16).\n"
                "양쪽에 유효한 점수가 있는 공통 시드만 사용하며 RUNNING/누락 점수는 제외합니다. "
                "update_tex.py의 EXCLUDED_SEEDS에 지정된 시드도 제외합니다. "
                "Drama/results/update_tex.py와 동일하게 표시된 시드 점수를 집계합니다."
            ),
            SCORE_DELTA_COLUMN: (
                "X행: Retrieval 미사용의 공통 시드 평균 − DramaXS 논문 점수.\n"
                "O행: |target: 16의 공통 시드 평균 − Retrieval 미사용의 공통 시드 평균|."
            ),
            HNS_DELTA_COLUMN: (
                "X행: (Retrieval 미사용 평균 − DramaXS 논문 점수) / (Human − Random).\n"
                "O행: (target: 16 평균 − Retrieval 미사용 평균) / (Human − Random).\n"
                "모든 평균은 EXCLUDED_SEEDS를 제외한 공통 시드 기준입니다. "
                "HNS는 Random=0, Human=1 기준이며 "
                "차이의 부호를 유지합니다. 백분율이 아닙니다. "
                "Random/Human은 STORM 엑셀과 동일한 기준값을 사용합니다."
            ),
        }
        for name, note in comments.items():
            results_sheet.cell(row=1, column=results.columns.get_loc(name) + 1).comment = Comment(note, "Drama")
        for row_index, game in enumerate(results["Game"], start=2):
            if results.at[row_index - 2, "Retrieval"] == "X" and game in REFERENCE_SCORES:
                random_score, human_score, _ = REFERENCE_SCORES[game]
                results_sheet.cell(
                    row=row_index, column=results.columns.get_loc(PAPER_SCORE_COLUMN) + 1,
                ).comment = Comment(f"Random: {random_score}\nHuman: {human_score}", "Drama")
        for row_index in range(len(running)):
            for column_index in range(len(seed_columns)):
                if running.iat[row_index, column_index] == "running":
                    cell = results_sheet.cell(row=row_index + 2, column=column_index + 3)
                    cell.value = "RUNNING"
                    cell.fill = running_fill
                    cell.font = Font(name="Calibri", size=12, bold=True, color="9C6500")
        mark_excluded_seeds(results_sheet, results, seed_columns, excluded_seeds)

        runs_sheet = writer.book["Runs"]
        for row_index, state in enumerate(detail["State"], start=2):
            if str(state).lower() == "running":
                for column_index in range(1, runs_sheet.max_column + 1):
                    runs_sheet.cell(row=row_index, column=column_index).fill = running_fill

        worksheet = writer.book["Results"]
        top_side = Side(style="medium", color="000000")
        previous_game = None
        for row_index in range(2, worksheet.max_row + 1):
            game = worksheet.cell(row=row_index, column=1).value
            if game != previous_game:
                for column_index in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=row_index, column=column_index)
                    cell.border = Border(
                        top=top_side,
                        left=cell.border.left,
                        right=cell.border.right,
                        bottom=cell.border.bottom,
                    )
                previous_game = game

    print(f"Successfully saved to {args.output}")


if __name__ == "__main__":
    main()
